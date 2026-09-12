"""report 子命令测试：取最近一次摸底 CSV、写入基线时间、生成 xlsx。"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import email_migration  # noqa: E402

BASE_HDR = ["email", "total_count", "total_size", "size_text(MB)"]
USER_HDR = ["EmailAddress", "Status", "SyncedItemCount", "SkippedItemCount",
            "EstimatedTotalTransferSize", "LastSuccessfulSyncTime"]
TGT_HDR = ["EmailAddress", "ItemCount", "TotalItemSize"]


def _write(path: Path, header: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        w.writerows(rows)


def _seed(out: Path) -> Path:
    d = out / "email-migration"
    _write(d / "email_check_success_20260910_000000.csv", BASE_HDR,
           [{"email": "a@src.example.com", "total_count": "10",
             "total_size": "1048576", "size_text(MB)": "1.00"}])
    _write(d / "email_check_success_20260911_000000.csv", BASE_HDR,
           [{"email": "a@src.example.com", "total_count": "20",
             "total_size": "2097152", "size_text(MB)": "2.00"}])
    _write(d / "user_stats.csv", USER_HDR,
           [{"EmailAddress": "a@corp.onmicrosoft.com", "Status": "Synced", "SyncedItemCount": "20",
             "SkippedItemCount": "0", "EstimatedTotalTransferSize": "2 MB (2097152 bytes)",
             "LastSuccessfulSyncTime": "2026/9/11 3:08:02"}])
    _write(d / "target_stats.csv", TGT_HDR,
           [{"EmailAddress": "a@corp.onmicrosoft.com", "ItemCount": "20",
             "TotalItemSize": "2 MB (2,097,152 bytes)"}])
    return d


def test_report_dry_run_picks_latest_baseline(tmp_path: Path, capsys) -> None:
    out = tmp_path / "out"
    _seed(out)
    assert email_migration.main(["report", "--output", str(out), "--dry-run"]) == 0
    printed = capsys.readouterr().out
    assert "email_check_success_20260911_000000.csv" in printed   # 取最新
    assert "email_check_success_20260910_000000.csv" not in printed
    assert "基线数据时间" not in printed  # dry-run 不生成报告


def test_report_writes_baseline_time_into_summary(tmp_path: Path) -> None:
    out = tmp_path / "out"
    _seed(out)
    xlsx = out / "rep.xlsx"
    assert email_migration.main(["report", "--output", str(out), "--out", str(xlsx),
                                 "--batch", "Migrate_Test_1"]) == 0
    wb = openpyxl.load_workbook(xlsx)
    assert "迁移数据汇总" in wb.sheetnames and "迁移概要" in wb.sheetnames
    summary = {(r[0].value, r[1].value) for r in wb["迁移概要"].iter_rows(min_row=1, max_row=30)
               if r[0].value}
    keys = {k for k, _ in summary}
    assert "基线文件" in keys and "基线数据时间" in keys and "更新时间" in keys
    assert dict(summary)["基线文件"] == "email_check_success_20260911_000000.csv"
    assert dict(summary)["迁移批次"] == "Migrate_Test_1"


def test_report_missing_inputs_returns_error(tmp_path: Path, capsys) -> None:
    out = tmp_path / "empty"
    (out / "email-migration").mkdir(parents=True)
    assert email_migration.main(["report", "--output", str(out), "--dry-run"]) == 1
    printed = capsys.readouterr().out
    assert "报告输入缺失" in printed


def test_report_warns_when_stats_older_than_baseline(tmp_path: Path, capsys) -> None:
    out = tmp_path / "out"
    d = _seed(out)
    old = 1_600_000_000
    # 把两个统计文件时间戳压到基线之前（mtime 设早）
    import os
    for name in ("user_stats.csv", "target_stats.csv"):
        os.utime(d / name, (old, old))
    assert email_migration.main(["report", "--output", str(out), "--dry-run"]) == 0
    assert "统计文件早于基线" in capsys.readouterr().out